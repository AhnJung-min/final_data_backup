"""Document extraction, image binding and immutable source evidence contracts."""
import io
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app.documents import DocumentCatalog, MAX_DOCUMENT_BYTES, parse_document
from app.inspector import inspect_text_rules
from app.main import create_app
from app.topics import DEFAULT_TOPIC, TopicCatalog

ROOT = Path(__file__).resolve().parents[1]


def template(topic=DEFAULT_TOPIC):
    return TopicCatalog(ROOT/'data/criteria.json', ROOT/'data/samples', ROOT/'data/topics').criteria(topic)


def parse(text, name='order.txt', topic=DEFAULT_TOPIC):
    return parse_document(text.encode('utf-8'), name, template(topic))


def test_explicit_conditions_do_not_fill_missing_values_from_demo():
    result = parse('발주 조건\n상품명: 말차라떼\n할인율: 20% 할인\n브랜드: OTHER BRAND\n추가조건: 세로형')
    assert [r['expected'] for r in result['rules']] == ['말차라떼', '20% 할인', 'OTHER BRAND']
    assert len(result['rules']) == 3
    assert result['rules'][0]['document_source'] == {'text': '상품명: 말차라떼', 'location': '2행'}
    assert result['document']['unparsed'][-1]['text'] == '추가조건: 세로형'
    assert len(result['warnings']) == 2
    assert parse('상품명: 딸기라떼')['version'] != result['version']


def test_document_brand_is_selected_by_layout_not_expected_words():
    rules = parse('브랜드: OTHER BRAND')['rules']
    def line(text):
        return {'text': text, 'bbox': [.2, .04, .5, .04], 'confidence': 1}
    assert inspect_text_rules([line('OTHER BRAND')], rules)[0]['verdict'] == '충족'
    changed = inspect_text_rules([line('DIFFERENT BRAND')], rules)[0]
    assert changed['observed'] == 'DIFFERENT BRAND'
    assert changed['verdict'] == '미충족'
    title = {'text': '딸기라떼', 'bbox': [.2, .065, .5, .13], 'confidence': 1}
    assert inspect_text_rules([line('OTHER BRAND'), title], rules)[0]['observed'] == 'OTHER BRAND'


@pytest.mark.parametrize('text,code', [
    ('아무 조건 없는 문서', 'NO_DOCUMENT_CONDITIONS'),
    ('상품명: 딸기라떼\n상품명: 말차라떼', 'DUPLICATE_CONDITION'),
    ('상품명:', 'EMPTY_CONDITION'),
])
def test_ambiguous_or_empty_conditions_are_actionable_errors(text, code):
    with pytest.raises(HTTPException) as error:
        parse(text)
    assert error.value.detail['code'] == code


def test_encodings_markdown_and_manual_rules():
    raw = '# 발주\n| 상품명 | 딸기라떼 |\n| 딸기 수량 | 4개 |'.encode('utf-16')
    result = parse_document(raw, 'order.md', template())
    assert [r['expected'] for r in result['rules']] == ['딸기라떼', '4개']
    rows = inspect_text_rules([], result['rules'])
    assert rows[1]['execution_status'] == 'not_run' and rows[1]['verdict'] == '확인 필요'


def test_docx_paragraphs_and_tables_retain_locations():
    from docx import Document
    document = Document()
    document.add_paragraph('상품명: 말차라떼')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '할인율'
    table.cell(0, 1).text = '20% 할인'
    output = io.BytesIO()
    document.save(output)
    result = parse_document(output.getvalue(), '발주.docx', template())
    assert [r['expected'] for r in result['rules']] == ['말차라떼', '20% 할인']
    assert result['rules'][1]['document_source']['location'] == '표 2 · 1행'


def pdf_bytes(text=None):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 20 250 Td ({text}) Tj ET'.encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_extracts_text_and_rejects_pages_without_text():
    result = parse_document(pdf_bytes('C01: Matcha Latte'), 'order.pdf', template())
    assert result['rules'][0]['expected'] == 'Matcha Latte'
    assert result['rules'][0]['document_source']['location'].startswith('1쪽')
    with pytest.raises(HTTPException) as error:
        parse_document(pdf_bytes(), 'scan.pdf', template())
    assert error.value.detail['code'] == 'PDF_TEXT_REQUIRED'


@pytest.mark.parametrize('raw,name,status', [(b'fake', 'fake.pdf', 422), (b'fake', 'fake.docx', 422),
                                           (b'', 'empty.txt', 422), (b'text', 'script.html', 415),
                                           (b'x'*(MAX_DOCUMENT_BYTES+1), 'large.txt', 413)],
                         ids=['fake-pdf', 'fake-docx', 'empty', 'unsupported', 'too-large'])
def test_rejects_invalid_or_oversized_uploads(raw, name, status):
    with pytest.raises(HTTPException) as error:
        parse_document(raw, name, template())
    assert error.value.status_code == status


@pytest.fixture
def client(tmp_path, monkeypatch):
    def inspect(path, rules):
        return {'results': inspect_text_rules([], rules), 'trace': [], 'method': 'test-only'}
    monkeypatch.setitem(sys.modules, 'app.mcp_client', SimpleNamespace(inspect_via_mcp=inspect))
    app = create_app(tmp_path / 'runtime')
    app.state.store.prepare()
    with TestClient(app, base_url='http://localhost') as client:
        client.headers['X-CSRFToken'] = client.get('/api/csrf').json()['csrf_token']
        yield client, tmp_path


def upload(client, text='상품명: 말차라떼\n할인율: 20% 할인'):
    response = client.post('/api/criteria/documents', data={'topic_id': DEFAULT_TOPIC},
                           files={'file': ('../발주.txt', text.encode('utf-8'), 'text/plain')})
    assert response.status_code == 200, response.text
    return response.json()


def test_document_confirmation_binding_source_history_and_restart(client):
    client, tmp_path = client
    document = upload(client)
    assert document['document']['name'] == '발주.txt' and not document['confirmed']
    asset = client.post('/api/assets/from-sample', json={'sample_id': 'sample-02',
                                                      'document_id': document['document_id']}).json()
    request = {'asset_id': asset['id'], 'criteria_version': document['version'], 'idempotency_key': 'doc-run'}
    assert client.post('/api/runs', json=request).status_code == 409
    response = client.post('/api/criteria/confirm', json={'version': document['version'],
                                                        'document_id': document['document_id']})
    assert response.status_code == 200 and response.json()['confirmed']
    different = upload(client, '상품명: 딸기라떼')
    assert client.post('/api/runs', json={**request, 'criteria_version': different['version']}).status_code == 409
    run = client.post('/api/runs', json=request).json()
    assert [r['expected'] for r in run['results']] == ['말차라떼', '20% 할인']
    assert run['results'][0]['document_source']['text'] == '상품명: 말차라떼'
    assert run['criteria_snapshot']['document']['name'] == '발주.txt'
    assert client.post('/api/runs', json=request).json()['id'] == run['id']
    with TestClient(create_app(tmp_path/'runtime'), base_url='http://localhost') as restarted:
        loaded = restarted.get('/api/criteria', params={'document_id': document['document_id']}).json()
        assert loaded['confirmed'] and loaded['rules'][0]['expected'] == '말차라떼'
        assert restarted.get('/api/runs/'+run['id']).json()['document_id'] == document['document_id']


def test_document_topic_and_path_isolation(client):
    client, _ = client
    document = upload(client)
    assert client.get('/api/criteria', params={'document_id': '../../secret'}).status_code == 404
    assert client.get('/api/criteria', params={'document_id': document['document_id'], 'topic_id': 'festival-v1'}).status_code == 409
    assert client.post('/api/assets/from-sample', json={'sample_id': 'sample-01', 'topic_id': 'festival-v1',
                                                      'document_id': document['document_id']}).status_code == 409
    client.headers.pop('X-CSRFToken')
    assert client.post('/api/criteria/documents', files={'file': ('test.txt', b'C01: Latte', 'text/plain')}).status_code == 403


def test_croissant_document_replaces_demo_order_counts(client):
    client, _ = client
    response = client.post('/api/criteria/documents', data={'topic_id': 'croissant-v2'},
                           files={'file': ('order.txt', '크루아상 수량: 큰 2개 · 미니 5개'.encode('utf-8'))})
    assert response.status_code == 200
    document = response.json()
    asset = client.post('/api/assets/from-sample', json={'sample_id': 'sample-01', 'topic_id': 'croissant-v2',
                                                      'document_id': document['document_id']}).json()
    assert asset['condition_id'] is None
    assert document['rules'][0]['expected'] == '큰 2개 · 미니 5개'
    assert 'expected_counts' not in document['rules'][0]


def test_template_is_a_downloadable_parseable_document(client):
    client, _ = client
    for topic in [DEFAULT_TOPIC, 'festival-v1', 'croissant-v2']:
        response = client.get('/api/criteria/template', params={'topic_id': topic})
        assert response.status_code == 200 and 'attachment' in response.headers['content-disposition']
        assert len(parse_document(response.content, 'template.txt', template(topic))['rules']) == 5
